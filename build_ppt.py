# -*- coding: utf-8 -*-
"""
AI Team Planner 해커톤 발표용 PowerPoint 생성 스크립트。
python-pptx 를 사용해 16:9 슬라이드를 작성한다.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- 색상 팔레트 ----
NAVY = RGBColor(0x1E, 0x2A, 0x4A)       # 진한 네이비 (제목/배경)
BLUE = RGBColor(0x2563, 0x2563 >> 8, 0xEB) if False else RGBColor(0x25, 0x63, 0xEB)  # 메인 블루
LIGHT_BLUE = RGBColor(0xDB, 0xEA, 0xFE)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
AMBER = RGBColor(0xD9, 0x77, 0x06)
RED = RGBColor(0xDC, 0x26, 0x26)
GRAY = RGBColor(0x6B, 0x72, 0x80)
DARK = RGBColor(0x1F, 0x29, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF3, 0xF4, 0xF6)

FONT = "Malgun Gothic"  # 한국어 대응 폰트 (Windows 기본)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width
SH = prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def fill_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_rect(slide, l, t, w, h, color, line_color=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def add_round_rect(slide, l, t, w, h, color, line_color=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


def add_text(slide, l, t, w, h, text, size=18, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.1):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return tb


def add_bullets(slide, l, t, w, h, items, size=16, color=DARK, gap=6, bold_lead=False):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
        # item 은 (bullet, rest) 형태의 튜플 또는 문자열
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run()
            r1.text = "● " + lead
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = BLUE
            r1.font.name = FONT
            r2 = p.add_run()
            r2.text = rest
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
            r2.font.name = FONT
        else:
            run = p.add_run()
            run.text = "• " + item
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.name = FONT
    return tb


def header_bar(slide, title, idx=None):
    """상단 제목 바."""
    add_rect(slide, 0, 0, SW, Inches(1.1), NAVY)
    add_rect(slide, 0, Inches(1.1), SW, Inches(0.06), BLUE)
    add_text(slide, Inches(0.5), 0, Inches(11.5), Inches(1.1), title,
             size=26, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if idx is not None:
        add_text(slide, Inches(12.3), 0, Inches(0.9), Inches(1.1), idx,
                 size=14, color=LIGHT_BLUE, bold=True,
                 align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


# =====================================================================
# 슬라이드 1 — 타이틀
# =====================================================================
s = add_slide()
fill_bg(s, NAVY)
add_rect(s, 0, Inches(3.3), SW, Inches(0.08), BLUE)
add_text(s, Inches(1), Inches(2.1), Inches(11.3), Inches(1.2),
         "AI Team Planner", size=54, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(3.5), Inches(11.3), Inches(0.8),
         "다음 주 출근·재택 예정 공유 & 출근 쏠림 자동 분석 플랫폼",
         size=22, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(5.6), Inches(11.3), Inches(1.0),
         "React + NestJS + PostgreSQL  |  해커톤 발표 자료",
         size=16, color=GRAY, align=PP_ALIGN.CENTER)

# =====================================================================
# 슬라이드 2 — 한 줄 소개 / 해결하는 문제
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "무엇을 만들었나", "01")
add_round_rect(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(1.5), LIGHT_BLUE)
add_text(s, Inches(0.9), Inches(1.6), Inches(11.5), Inches(1.3),
         "“다음 주 팀원들의 출근/재택 예정을 공유하고,\n관리자가 출근 인원의 쏠림(과다·과소)을 자동 감지해 오피스 가동을 최적화한다.”",
         size=20, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)

add_text(s, Inches(0.6), Inches(3.3), Inches(6), Inches(0.5), "해결하려는 문제", size=20, color=BLUE, bold=True)
add_bullets(s, Inches(0.7), Inches(3.9), Inches(5.9), Inches(3),
            ["하이브리드 근무로 '누가 언제 출근하는지' 파악이 어려움",
             "특정 요일에 자리가 몰리거나(혼잡), 텅 비어(비효율) 공간 낭비",
             "관리자가 출근 현황을 수작업으로 집계"], size=16, gap=12)

add_text(s, Inches(6.8), Inches(3.3), Inches(6), Inches(0.5), "제공하는 가치", size=20, color=GREEN, bold=True)
add_bullets(s, Inches(6.9), Inches(3.9), Inches(5.8), Inches(3),
            ["팀 캘린더로 한눈에 주간 출근 현황 공유",
             "출근 상한/하한 임계치 기반 과밀·과소 자동 경고",
             "출근율·출근자 명단을 관리 대시보드로 가시화"], size=16, gap=12)

# =====================================================================
# 슬라이드 3 — 시스템 아키텍처
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "시스템 아키텍처", "02")

# 3개 박스 (브라우저 → 백엔드 → DB)
box_y = Inches(2.2)
box_h = Inches(1.6)
w = Inches(3.5)
gap_x = Inches(0.85)
x0 = Inches(0.75)

def arch_box(x, title, sub, color):
    add_round_rect(s, x, box_y, w, box_h, color)
    add_text(s, x, box_y + Inches(0.25), w, Inches(0.6), title, size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x, box_y + Inches(0.85), w, Inches(0.6), sub, size=13, color=WHITE, align=PP_ALIGN.CENTER)

arch_box(x0, "프론트엔드", "React 18 + Vite + TS\nTailwind CSS  (:5173)", BLUE)
arch_box(x0 + w + gap_x, "백엔드 API", "NestJS + TypeScript\n모든 경로 /api  (:3000)", NAVY)
arch_box(x0 + 2*(w + gap_x), "데이터베이스", "PostgreSQL + TypeORM\n(:5432)", GREEN)

# 화살표 텍스트
add_text(s, x0 + w, box_y + Inches(0.1), gap_x, Inches(0.5), "→", size=28, color=GRAY, align=PP_ALIGN.CENTER)
add_text(s, x0 + 2*w + gap_x, box_y + Inches(0.1), gap_x, Inches(0.5), "→", size=28, color=GRAY, align=PP_ALIGN.CENTER)
add_text(s, x0, box_y - Inches(0.55), Inches(12), Inches(0.5),
         "fetch + Authorization: Bearer <JWT> 자동 첨부            TypeORM 쿼리", size=12, color=GRAY, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.6), Inches(4.4), Inches(12), Inches(0.5), "설계 특징", size=20, color=BLUE, bold=True)
add_bullets(s, Inches(0.7), Inches(5.0), Inches(12), Inches(2),
            [("3계층 분리: ", "얇은 Controller → Service(DB/정형화) → 순수 도메인 함수(부작용 없음)"),
             ("Docker Compose: ", "frontend / backend / postgres 3개 컨테이너를 한 번에 기동"),
             ("UTC 고정 날짜 계산: ", "타임존 off-by-one 버그를 원천 차단, 결정론적 결과")], size=15, gap=9)

# =====================================================================
# 슬라이드 4 — 기술 스택
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "기술 스택", "03")
cols = [
    ("프론트엔드", BLUE, ["React 18", "Vite", "TypeScript", "TailwindCSS", "React Router v6"]),
    ("백엔드", NAVY, ["NestJS 10", "TypeScript", "TypeORM", "class-validator", "JWT (@nestjs/jwt)"]),
    ("데이터 / 인증", GREEN, ["PostgreSQL 16", "bcrypt (해싱)", "자체 발행 JWT", "Docker Compose", "Jest (테스트)"]),
]
cw = Inches(3.9)
cx0 = Inches(0.6)
cgap = Inches(0.25)
for i, (title, color, items) in enumerate(cols):
    x = cx0 + i * (cw + cgap)
    add_round_rect(s, x, Inches(1.6), cw, Inches(0.8), color)
    add_text(s, x, Inches(1.6), cw, Inches(0.8), title, size=18, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_round_rect(s, x, Inches(2.5), cw, Inches(3.6), LIGHT_BG)
    add_bullets(s, x + Inches(0.3), Inches(2.75), cw - Inches(0.5), Inches(3.2), items, size=16, gap=14)

# =====================================================================
# 슬라이드 5 — 권한 체계 (3단계)
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "권한 체계 — 역할 3단계", "04")
roles = [
    ("employee (직원)", GRAY, "본인 근무 예정 등록/조회\n팀 캘린더 보기"),
    ("team_admin (팀 관리자)", AMBER, "직원 기능 +\n대시보드/AI 분석 조회\n멤버 팀 배정"),
    ("super_admin (최상위 관리자)", BLUE, "모든 기능 +\n역할 변경 · 비밀번호 재설정\n팀 생성/삭제"),
]
rw = Inches(3.9)
rx0 = Inches(0.6)
rgap = Inches(0.25)
for i, (title, color, desc) in enumerate(roles):
    x = rx0 + i * (rw + rgap)
    add_round_rect(s, x, Inches(1.7), rw, Inches(3.4), WHITE, line_color=color)
    add_rect(s, x, Inches(1.7), rw, Inches(0.12), color)
    add_text(s, x, Inches(2.0), rw, Inches(0.8), title, size=16, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.9), rw - Inches(0.4), Inches(2.0), desc, size=14, color=DARK, align=PP_ALIGN.CENTER)

add_round_rect(s, Inches(0.6), Inches(5.4), Inches(12.1), Inches(1.4), LIGHT_BLUE)
add_text(s, Inches(0.9), Inches(5.55), Inches(11.5), Inches(1.1),
         "이중 권한 제어 — 서버(JwtAuthGuard → RolesGuard)와 클라이언트(ProtectedRoute)\n양쪽에서 역할을 검증. 메뉴 자체도 역할에 따라 노출/숨김 처리.",
         size=15, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 슬라이드 6 — 핵심 개념 (Target_Week / Occupancy_Count)
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "핵심 도메인 개념", "05")
add_round_rect(s, Inches(0.6), Inches(1.7), Inches(5.9), Inches(4.5), LIGHT_BG)
add_rect(s, Inches(0.6), Inches(1.7), Inches(5.9), Inches(0.9), BLUE)
add_text(s, Inches(0.6), Inches(1.7), Inches(5.9), Inches(0.9), "Target_Week (타깃 주)",
         size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_bullets(s, Inches(0.9), Inches(2.9), Inches(5.3), Inches(3),
            ["기준일로부터 '다음 주 월~금' 평일 5일",
             "토·일은 항상 제외",
             "모든 계산을 UTC로 고정 → 타임존 버그 방지",
             "예정 등록·캘린더·분석 모두 이 주간을 대상"], size=15, gap=12)

add_round_rect(s, Inches(6.8), Inches(1.7), Inches(5.9), Inches(4.5), LIGHT_BG)
add_rect(s, Inches(6.8), Inches(1.7), Inches(5.9), Inches(0.9), GREEN)
add_text(s, Inches(6.8), Inches(1.7), Inches(5.9), Inches(0.9), "Occupancy_Count (출근 인원 수)",
         size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_bullets(s, Inches(7.1), Inches(2.9), Inches(5.3), Inches(3),
            ["특정 날짜에 'office(출근)'로 등록한 사람 수",
             "평균이 아니라 '카운트(개수)'",
             "대시보드·AI 분석의 모든 판정 기준",
             "출근율 = officeCount / 전체등록자 × 100"], size=15, gap=12)

# =====================================================================
# 슬라이드 7 — 페이지 맵 (8개 화면)
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "화면 구성 — 8개 페이지", "06")
pages = [
    ("① 로그인 /login", "이메일/비밀번호 인증", GRAY),
    ("② 회원가입 /register", "가입 후 자동 로그인", GRAY),
    ("③ 팀 캘린더 /", "주간 출근 현황 표 (전체)", BLUE),
    ("④ 근무 예정 /schedule", "다음 주 출근/재택 즉시 등록", BLUE),
    ("⑤ 대시보드 /dashboard", "출근율·명단·임계치 (관리자)", AMBER),
    ("⑥ AI 분석 /analysis ★", "과다/과소 자동 경고 (관리자)", RED),
    ("⑦ 사용자 관리 /admin/users", "팀 배정·역할·비번 (관리자)", AMBER),
    ("⑧ 팀 관리 /admin/teams", "팀 생성/삭제 (최상위)", BLUE),
]
pw = Inches(5.95)
ph = Inches(1.15)
px0 = Inches(0.6)
py0 = Inches(1.5)
pgx = Inches(0.25)
pgy = Inches(0.2)
for i, (title, desc, color) in enumerate(pages):
    col = i % 2
    row = i // 2
    x = px0 + col * (pw + pgx)
    y = py0 + row * (ph + pgy)
    add_round_rect(s, x, y, pw, ph, WHITE, line_color=color)
    add_rect(s, x, y, Inches(0.12), ph, color)
    add_text(s, x + Inches(0.3), y + Inches(0.12), pw - Inches(0.4), Inches(0.5), title, size=15, color=NAVY, bold=True)
    add_text(s, x + Inches(0.3), y + Inches(0.58), pw - Inches(0.4), Inches(0.5), desc, size=13, color=GRAY)

# =====================================================================
# 슬라이드 8 — 사용자 흐름
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "사용자 흐름 (데모 시나리오)", "07")

add_text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(0.5), "직원 (Employee)", size=18, color=BLUE, bold=True)
flow1 = ["로그인", "근무 예정 등록 (/schedule)", "팀 캘린더에서 전체 확인 (/)"]
fx = Inches(0.7)
for i, step in enumerate(flow1):
    bw = Inches(3.4)
    add_round_rect(s, fx, Inches(2.1), bw, Inches(0.9), LIGHT_BLUE)
    add_text(s, fx, Inches(2.1), bw, Inches(0.9), step, size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(flow1) - 1:
        add_text(s, fx + bw, Inches(2.1), Inches(0.5), Inches(0.9), "→", size=24, color=GRAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    fx = fx + bw + Inches(0.5)

add_text(s, Inches(0.6), Inches(3.5), Inches(12), Inches(0.5), "관리자 (Admin)", size=18, color=AMBER, bold=True)
flow2 = ["로그인", "대시보드 확인·임계치 설정", "AI 분석 실행 → 경고 확인", "팀·사용자 관리"]
fx = Inches(0.7)
for i, step in enumerate(flow2):
    bw = Inches(2.7)
    add_round_rect(s, fx, Inches(4.1), bw, Inches(0.9), RGBColor(0xFE, 0xF3, 0xC7))
    add_text(s, fx, Inches(4.1), bw, Inches(0.9), step, size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(flow2) - 1:
        add_text(s, fx + bw, Inches(4.1), Inches(0.35), Inches(0.9), "→", size=20, color=GRAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    fx = fx + bw + Inches(0.35)

add_round_rect(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.1), LIGHT_BG)
add_text(s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(0.9),
         "기본 관리자 계정(시드): admin@example.com / Passw0rd!  (super_admin)",
         size=15, color=DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 슬라이드 9 — AI 분석 로직 (핵심)
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "AI 분석 — 어떻게 동작하나 ★", "08")
add_round_rect(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.8), NAVY)
add_text(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.8),
         "임계치 기반 규칙 엔진 — 관리자가 정한 상한/하한으로 출근 쏠림을 자동 감지",
         size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

add_text(s, Inches(0.6), Inches(2.5), Inches(12), Inches(0.5), "실행 4단계", size=18, color=BLUE, bold=True)
steps = [
    ("1. 대상 주간 결정", "서버 오늘(UTC) 기준 다음 주 월~금 5일 계산"),
    ("2. 근무 예정 조회", "조직 전체 Schedule 레코드를 읽기 전용으로 조회"),
    ("3. 임계치 로드", "관리자 설정 상한(upper)·하한(lower) 가져오기 (기본 70/30)"),
    ("4. 판정", "각 날짜 출근자 수를 상한/하한과 비교해 경고 생성"),
]
sy = Inches(3.05)
for i, (t, d) in enumerate(steps):
    add_round_rect(s, Inches(0.7), sy, Inches(5.9), Inches(0.85), LIGHT_BG)
    add_text(s, Inches(0.9), sy + Inches(0.08), Inches(5.6), Inches(0.4), t, size=14, color=NAVY, bold=True)
    add_text(s, Inches(0.9), sy + Inches(0.45), Inches(5.6), Inches(0.4), d, size=11.5, color=GRAY)
    sy = sy + Inches(0.95)

# 판정 규칙 코드 박스
add_round_rect(s, Inches(6.9), Inches(3.05), Inches(5.8), Inches(3.65), DARK)
code = ("각 날짜(월~금):\n"
        "  occupancyCount = 그날 출근자 수\n\n"
        "  if occupancyCount >= upper:\n"
        "      → 출근 과다 경고 (over)\n"
        "  if occupancyCount <= lower:\n"
        "      → 출근 과소 경고 (under)\n\n"
        "  요약표에 {출근, 재택} 추가")
add_text(s, Inches(7.2), Inches(3.3), Inches(5.3), Inches(3.2), code,
         size=14, color=RGBColor(0x8B, 0xE9, 0xA0), font="Consolas", line_spacing=1.15)

# =====================================================================
# 슬라이드 10 — AI 분석 Q&A (솔직한 설명)
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "AI 분석 — 예상 질문 대응", "09")
qa = [
    ("Q. 머신러닝 모델인가요?", "아니요. 규칙 기반 조건 판정(if)입니다. 블랙박스가 아닌 투명·재현 가능한 방식.", RED),
    ("Q. 그냥 평균 아닌가요?", "아니요. 출근자 '수를 세어(count)' 관리자가 정한 상한/하한과 '비교'합니다.", AMBER),
    ("Q. 왜 즉시 결과가 나오나요?", "5일치 가벼운 비교 연산뿐. 모델 추론·외부 호출이 없어 밀리초 단위 응답.", BLUE),
    ("Q. 그럼 가치가 있나요?", "결과가 투명·설명 가능. 판정부만 교체하면 ML 예측으로 확장 가능한 구조.", GREEN),
]
qy = Inches(1.6)
for q, a, color in qa:
    add_round_rect(s, Inches(0.6), qy, Inches(12.1), Inches(1.15), WHITE, line_color=color)
    add_rect(s, Inches(0.6), qy, Inches(0.12), Inches(1.15), color)
    add_text(s, Inches(0.95), qy + Inches(0.12), Inches(11.5), Inches(0.5), q, size=16, color=NAVY, bold=True)
    add_text(s, Inches(0.95), qy + Inches(0.6), Inches(11.5), Inches(0.5), a, size=14, color=DARK)
    qy = qy + Inches(1.3)

# =====================================================================
# 슬라이드 11 — 데이터 모델
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "데이터 모델 (테이블 4개)", "10")
entities = [
    ("User", BLUE, "id, email, password_hash,\nname, role, team_id\n→ Team에 N:1"),
    ("Team", NAVY, "id, name, leader_id\n→ User를 1:N 보유"),
    ("Schedule", GREEN, "id, user_id, date,\nwork_location(office/remote)\n(user_id,date) 복합 유니크"),
    ("ThresholdSetting", AMBER, "upper_threshold,\nlower_threshold, updated_by\n시스템 전체 단일 레코드"),
]
ew = Inches(2.95)
ex0 = Inches(0.55)
egap = Inches(0.18)
for i, (name, color, desc) in enumerate(entities):
    x = ex0 + i * (ew + egap)
    add_round_rect(s, x, Inches(1.8), ew, Inches(0.8), color)
    add_text(s, x, Inches(1.8), ew, Inches(0.8), name, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_round_rect(s, x, Inches(2.7), ew, Inches(2.6), LIGHT_BG)
    add_text(s, x + Inches(0.2), Inches(2.95), ew - Inches(0.4), Inches(2.2), desc, size=13, color=DARK, align=PP_ALIGN.CENTER)

add_round_rect(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.1), LIGHT_BLUE)
add_text(s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(0.9),
         "(user_id, date) 복합 유니크 제약 → '한 사람이 한 날짜에 1건'을 DB가 보장 (upsert의 근거)",
         size=14, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 슬라이드 12 — 설계 어필 포인트
# =====================================================================
s = add_slide()
fill_bg(s, WHITE)
header_bar(s, "설계상의 강점", "11")
points = [
    ("3계층 분리 ", "— Controller / Service / 순수 도메인 함수로 로직 중복 제거, 테스트 용이"),
    ("단일 정보원(SSOT) ", "— 역할·근무구분 enum, 임계치, 집계 로직을 한 곳에 집약"),
    ("결정론적 날짜 계산 ", "— 모든 날짜를 UTC로 고정, 타임존 off-by-one 버그 차단"),
    ("분석 읽기 전용 보장 ", "— 몇 번을 실행해도 데이터가 변하지 않아 안전"),
    ("이중 권한 제어 ", "— 서버 가드 + 클라이언트 라우트 양쪽에서 역할 검증"),
    ("보안 기본기 ", "— bcrypt 해싱, 로그인 실패 시 계정 추측 방지 동일 메시지, 자체 JWT"),
]
add_bullets(s, Inches(0.8), Inches(1.7), Inches(11.8), Inches(5),
            points, size=18, gap=18)

# =====================================================================
# 슬라이드 13 — 마무리
# =====================================================================
s = add_slide()
fill_bg(s, NAVY)
add_rect(s, 0, Inches(3.0), SW, Inches(0.08), BLUE)
add_text(s, Inches(1), Inches(2.4), Inches(11.3), Inches(0.8),
         "감사합니다", size=44, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(3.4), Inches(11.3), Inches(1.0),
         "AI Team Planner — 출근 쏠림을 미리 보고, 공간을 똑똑하게",
         size=18, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

prs.save("AI_Team_Planner_발표자료.pptx")
print("saved: AI_Team_Planner_발표자료.pptx  / slides:", len(prs.slides._sldIdLst))
